# -*- coding: utf-8 -*-
"""
修改前3章PPT：
1. 给每章前3页加思考题
2. 在第1章PPT中插入NIST 3-5-4渐进展示幻灯片
"""
import sys, os, copy
sys.stdout.reconfigure(encoding='utf-8')

from pptx import Presentation
from pptx.util import Inches, Pt, Emu, Cm
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

SRC = r'E:\xiqin\云计算基础技术与应用-PPT课件.rar\53938-云计算基础技术与应用-PPT课件'
OUT = r'D:\WorkBuddyData\2026-08-11-07-48-22\modified_ppt'
os.makedirs(OUT, exist_ok=True)

# ============================================================
# 颜色定义
# ============================================================
C_BLUE   = RGBColor(0x00, 0x78, 0xD4)   # 微软蓝
C_DARK   = RGBColor(0x1F, 0x38, 0x64)   # 深蓝
C_WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
C_ORANGE = RGBColor(0xF3, 0x8B, 0x00)   # 橙色（问题标签）
C_GREEN  = RGBColor(0x10, 0x7C, 0x10)   # 绿色
C_RED    = RGBColor(0xC4, 0x31, 0x27)   # 红色
C_PURPLE = RGBColor(0x5C, 0x2D, 0x91)   # 紫色
C_TEAL   = RGBColor(0x00, 0x80, 0x80)   # 青色
C_LIGHT_BLUE = RGBColor(0xE8, 0xF1, 0xFA)
C_LIGHT_ORANGE = RGBColor(0xFD, 0xF1, 0xE0)
C_LIGHT_GREEN = RGBColor(0xE8, 0xF5, 0xE9)
C_LIGHT_RED   = RGBColor(0xFD, 0xEB, 0xE8)
C_LIGHT_PURPLE = RGBColor(0xF3, 0xE8, 0xFA)
C_LIGHT_TEAL  = RGBColor(0xE0, 0xF5, 0xF5)
C_GRAY_BG = RGBColor(0xF5, 0xF5, 0xF5)
C_DARK_GRAY = RGBColor(0x40, 0x40, 0x40)

# 幻灯片尺寸
SLIDE_W = 12192000  # EMU
SLIDE_H = 6858000   # EMU

# ============================================================
# 工具函数
# ============================================================
def add_question_box(slide, question_text, top=6.0):
    """在幻灯片底部加一个橙色问题文本框"""
    left = Inches(0.5)
    top_emu = Inches(top)
    width = Inches(12.3)
    height = Inches(1.1)
    
    shape = slide.shapes.add_textbox(left, top_emu, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = C_LIGHT_ORANGE
    shape.line.color.rgb = C_ORANGE
    shape.line.width = Pt(1.5)
    
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Pt(12)
    tf.margin_right = Pt(12)
    tf.margin_top = Pt(6)
    tf.margin_bottom = Pt(6)
    
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    run = p.add_run()
    run.text = "思考题  "
    run.font.size = Pt(16)
    run.font.bold = True
    run.font.color.rgb = C_ORANGE
    
    run2 = p.add_run()
    run2.text = question_text
    run2.font.size = Pt(16)
    run2.font.color.rgb = C_DARK_GRAY
    run2.font.name = '微软雅黑'

def add_textbox(slide, text, left, top, width, height, font_size=18, bold=False, color=C_DARK, bg_color=None, align=PP_ALIGN.LEFT, font_name='微软雅黑'):
    """通用文本框"""
    shape = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    if bg_color:
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_color
    else:
        shape.fill.background()
    shape.line.fill.background()
    
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Pt(8)
    tf.margin_right = Pt(8)
    tf.margin_top = Pt(4)
    tf.margin_bottom = Pt(4)
    
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font_name
    return shape

def add_rounded_rect(slide, left, top, width, height, fill_color, line_color=None):
    """加圆角矩形"""
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(left), Inches(top), Inches(width), Inches(height)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if line_color:
        shape.line.color.rgb = line_color
        shape.line.width = Pt(2)
    else:
        shape.line.fill.background()
    return shape

def add_rect(slide, left, top, width, height, fill_color, line_color=None):
    """加矩形"""
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(left), Inches(top), Inches(width), Inches(height)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if line_color:
        shape.line.color.rgb = line_color
        shape.line.width = Pt(2)
    else:
        shape.line.fill.background()
    return shape

def add_shape_with_text(slide, left, top, width, height, text, fill_color, text_color=None, font_size=16, bold=True, shape_type=MSO_SHAPE.ROUNDED_RECTANGLE):
    """加带文字的形状"""
    shape = slide.shapes.add_shape(
        shape_type,
        Inches(left), Inches(top), Inches(width), Inches(height)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    shape.line.fill.background()
    
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Pt(6)
    tf.margin_right = Pt(6)
    tf.margin_top = Pt(4)
    tf.margin_bottom = Pt(4)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = text
    run.font.size = Pt(font_size)
    run.font.bold = bold
    if text_color:
        run.font.color.rgb = text_color
    else:
        run.font.color.rgb = C_WHITE
    run.font.name = '微软雅黑'
    return shape

def add_arrow(slide, left, top, width, height):
    """加箭头"""
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RIGHT_ARROW,
        Inches(left), Inches(top), Inches(width), Inches(height)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = C_DARK_GRAY
    shape.line.fill.background()
    return shape

# ============================================================
# 第1步：给3个PPT的前3页加思考题
# ============================================================
questions = {
    '1.云计算的简介.pptx': [
        '你手机里的照片存在哪里？微信消息存在哪里？它们在\u201c云端\u201d吗？',
        '学完这章你能回答：云计算到底是什么？它和我们用的网盘有什么区别？',
        '你心中的\u201c云\u201d是什么？说说你生活中用到的\u201c云\u201d服务。',
    ],
    '2.云计算的服务.pptx': [
        '上节课学的NIST 3-5-4框架中的\u201c3\u201d指的是什么？你能说出几个？',
        '你用过的软件，哪些是\u201c装在电脑上\u201d的，哪些是\u201c在网页上直接用\u201d的？',
        'SaaS、PaaS、IaaS有什么区别？带着这个问题看本章目录。',
    ],
    '3.云计算的部署.pptx': [
        '上节课学的NIST 3-5-4框架中的\u201c4\u201d指的是什么？',
        '学校的机房和阿里云，哪个更省钱？为什么大企业还要自建机房？',
        '公有云、私有云、混合云有什么区别？带着这个问题看本章目录。',
    ],
}

print("=" * 60)
print("第1步：给3个PPT的前3页加思考题")
print("=" * 60)

for fname, qs in questions.items():
    src_path = os.path.join(SRC, fname)
    out_path = os.path.join(OUT, fname)
    prs = Presentation(src_path)
    
    for i, slide in enumerate(prs.slides):
        if i >= 3:
            break
        add_question_box(slide, qs[i])
    
    prs.save(out_path)
    print(f"  [OK] {fname} -> {out_path}")

# ============================================================
# 第2步：在第1章PPT中插入NIST 3-5-4渐进展示幻灯片
# ============================================================
print()
print("=" * 60)
print("第2步：在第1章PPT插入NIST 3-5-4渐进展示幻灯片")
print("=" * 60)

prs1 = Presentation(os.path.join(OUT, '1.云计算的简介.pptx'))

# 获取空白布局
blank_layout = prs1.slide_layouts[6]  # 通常6是空白布局

def make_nist_slide_1(prs):
    """幻灯片1：NIST是什么？"""
    slide = prs.slides.add_slide(blank_layout)
    # 背景
    bg = add_rect(slide, 0, 0, 13.3, 7.5, C_DARK)
    # 标题
    add_textbox(slide, "NIST 是什么？", 1.5, 0.5, 10.3, 1.0, font_size=36, bold=True, color=C_WHITE, align=PP_ALIGN.CENTER)
    # 分隔线
    add_rect(slide, 3, 1.5, 7.3, 0.06, C_ORANGE)
    # 内容
    add_textbox(slide, "NIST = National Institute of Standards and Technology", 1.5, 2.0, 10.3, 0.8, font_size=22, bold=True, color=C_ORANGE, align=PP_ALIGN.CENTER)
    add_textbox(slide, "美国国家标准与技术研究院", 1.5, 2.7, 10.3, 0.6, font_size=20, color=C_LIGHT_BLUE, align=PP_ALIGN.CENTER)
    
    # 三个要点
    items = [
        ("成立于1901年", "隶属美国商务部，是全球权威的标准化机构"),
        ("SP 800-145文件", "2011年9月发布云计算定义，被全球广泛引用"),
        ("为什么重要？", "NIST的定义是国际上最权威、被引用最多的云计算定义"),
    ]
    for i, (title, desc) in enumerate(items):
        y = 3.6 + i * 1.15
        add_rounded_rect(slide, 1.5, y, 10.3, 0.95, C_DARK_GRAY)
        add_textbox(slide, title, 1.8, y + 0.05, 2.8, 0.85, font_size=16, bold=True, color=C_ORANGE)
        add_textbox(slide, desc, 4.5, y + 0.05, 7.0, 0.85, font_size=15, color=C_WHITE)
    
    return slide

def make_nist_slide_2(prs):
    """幻灯片2：NIST云计算定义"""
    slide = prs.slides.add_slide(blank_layout)
    add_rect(slide, 0, 0, 13.3, 7.5, C_WHITE)
    # 标题栏
    add_rect(slide, 0, 0, 13.3, 1.2, C_DARK)
    add_textbox(slide, "NIST 云计算定义", 0.5, 0.2, 12.3, 0.8, font_size=30, bold=True, color=C_WHITE, align=PP_ALIGN.CENTER)
    
    # 英文定义框
    add_rounded_rect(slide, 0.5, 1.5, 12.3, 2.0, C_LIGHT_BLUE)
    add_textbox(slide, "Cloud computing is a model for enabling ubiquitous, convenient, on-demand network access to a shared pool of configurable computing resources that can be rapidly provisioned and released with minimal management effort or service provider interaction.", 
                0.8, 1.6, 11.7, 1.8, font_size=14, color=C_DARK_GRAY)
    add_textbox(slide, "—— NIST SP 800-145 (2011.09)", 8.5, 3.2, 4.0, 0.4, font_size=12, color=C_ORANGE, align=PP_ALIGN.RIGHT)
    
    # 中文翻译框
    add_rounded_rect(slide, 0.5, 3.8, 12.3, 2.0, C_LIGHT_GREEN)
    add_textbox(slide, "云计算是一种模型，它允许通过无处不在的、便捷的、按需的网络访问，来获取一个可配置的计算资源共享池（如网络、服务器、存储、应用和服务），这些资源能够被快速分配和释放，且只需最少的管理精力或服务提供商交互。",
                0.8, 3.9, 11.7, 1.8, font_size=15, color=C_DARK_GRAY)
    
    # 底部提示
    add_rounded_rect(slide, 3, 6.1, 7.3, 0.9, C_ORANGE)
    add_textbox(slide, "这个定义 = 3个服务模型 + 5个基本特征 + 4个部署模型", 3.2, 6.15, 6.9, 0.8, font_size=16, bold=True, color=C_WHITE, align=PP_ALIGN.CENTER)
    
    return slide

def make_nist_slide_3(prs):
    """幻灯片3：3-5-4框架总览"""
    slide = prs.slides.add_slide(blank_layout)
    add_rect(slide, 0, 0, 13.3, 7.5, C_WHITE)
    # 标题栏
    add_rect(slide, 0, 0, 13.3, 1.2, C_DARK)
    add_textbox(slide, "NIST 3-5-4 框架总览", 0.5, 0.2, 12.3, 0.8, font_size=30, bold=True, color=C_WHITE, align=PP_ALIGN.CENTER)
    add_textbox(slide, "记住 3-5-4，就记住了云计算的全貌", 3, 0.85, 7.3, 0.4, font_size=14, color=C_LIGHT_BLUE, align=PP_ALIGN.CENTER)
    
    # 三个大圆/方块
    # 3 - 服务模型
    add_shape_with_text(slide, 1.0, 2.0, 3.3, 3.5, "3", C_BLUE, font_size=72, bold=True)
    add_textbox(slide, "服务模型", 1.0, 5.0, 3.3, 0.5, font_size=20, bold=True, color=C_BLUE, align=PP_ALIGN.CENTER)
    add_textbox(slide, "Service Models", 1.0, 5.4, 3.3, 0.4, font_size=14, color=C_DARK_GRAY, align=PP_ALIGN.CENTER)
    add_textbox(slide, "SaaS / PaaS / IaaS", 1.0, 5.8, 3.3, 0.4, font_size=13, color=C_DARK_GRAY, align=PP_ALIGN.CENTER)
    
    # 5 - 基本特征
    add_shape_with_text(slide, 5.0, 2.0, 3.3, 3.5, "5", C_GREEN, font_size=72, bold=True)
    add_textbox(slide, "基本特征", 5.0, 5.0, 3.3, 0.5, font_size=20, bold=True, color=C_GREEN, align=PP_ALIGN.CENTER)
    add_textbox(slide, "Essential Characteristics", 5.0, 5.4, 3.3, 0.4, font_size=14, color=C_DARK_GRAY, align=PP_ALIGN.CENTER)
    add_textbox(slide, "按需自助 / 网络访问 / 资源池化 / 快速弹性 / 可计量", 4.5, 5.8, 4.3, 0.4, font_size=12, color=C_DARK_GRAY, align=PP_ALIGN.CENTER)
    
    # 4 - 部署模型
    add_shape_with_text(slide, 9.0, 2.0, 3.3, 3.5, "4", C_ORANGE, font_size=72, bold=True)
    add_textbox(slide, "部署模型", 9.0, 5.0, 3.3, 0.5, font_size=20, bold=True, color=C_ORANGE, align=PP_ALIGN.CENTER)
    add_textbox(slide, "Deployment Models", 9.0, 5.4, 3.3, 0.4, font_size=14, color=C_DARK_GRAY, align=PP_ALIGN.CENTER)
    add_textbox(slide, "公有云 / 私有云 / 社区云 / 混合云", 8.8, 5.8, 3.7, 0.4, font_size=12, color=C_DARK_GRAY, align=PP_ALIGN.CENTER)
    
    # 底部提示
    add_textbox(slide, "接下来逐一拆解 →", 9.5, 6.6, 3.3, 0.5, font_size=14, bold=True, color=C_ORANGE, align=PP_ALIGN.RIGHT)
    
    return slide

def make_nist_slide_4(prs):
    """幻灯片4：3个服务模型"""
    slide = prs.slides.add_slide(blank_layout)
    add_rect(slide, 0, 0, 13.3, 7.5, C_WHITE)
    # 标题栏
    add_rect(slide, 0, 0, 13.3, 1.2, C_BLUE)
    add_textbox(slide, "3-5-4 之【3】三个服务模型", 0.5, 0.2, 12.3, 0.8, font_size=28, bold=True, color=C_WHITE, align=PP_ALIGN.CENTER)
    
    # 三列
    models = [
        ("SaaS", "软件即服务", "Software as a Service", "用户直接使用云端应用", "在线Office、钉钉、企业微信", C_BLUE, C_LIGHT_BLUE),
        ("PaaS", "平台即服务", "Platform as a Service", "用户在云端开发部署应用", "Gitee代码仓库、微信小程序云开发", C_GREEN, C_LIGHT_GREEN),
        ("IaaS", "基础设施即服务", "Infrastructure as a Service", "用户租用虚拟机/存储/网络", "阿里云ECS、腾讯云CVM", C_ORANGE, C_LIGHT_ORANGE),
    ]
    
    for i, (abbr, cn, en, desc, example, color, bg) in enumerate(models):
        x = 0.5 + i * 4.2
        # 标题块
        add_shape_with_text(slide, x, 1.5, 3.8, 1.0, abbr, color, font_size=28)
        # 中文名
        add_textbox(slide, cn, x, 2.6, 3.8, 0.5, font_size=18, bold=True, color=color, align=PP_ALIGN.CENTER)
        # 英文名
        add_textbox(slide, en, x, 3.1, 3.8, 0.4, font_size=12, color=C_DARK_GRAY, align=PP_ALIGN.CENTER)
        # 描述
        add_rounded_rect(slide, x, 3.7, 3.8, 1.0, bg)
        add_textbox(slide, desc, x + 0.1, 3.75, 3.6, 0.9, font_size=14, color=C_DARK_GRAY, align=PP_ALIGN.CENTER)
        # 举例
        add_textbox(slide, "举例：", x, 4.8, 3.8, 0.4, font_size=13, bold=True, color=color, align=PP_ALIGN.CENTER)
        add_textbox(slide, example, x, 5.2, 3.8, 0.7, font_size=13, color=C_DARK_GRAY, align=PP_ALIGN.CENTER)
    
    # 底部类比
    add_rounded_rect(slide, 0.5, 6.1, 12.3, 1.0, C_LIGHT_ORANGE)
    add_textbox(slide, "生活类比：  SaaS = 去饭店吃饭（菜做好直接吃）   PaaS = 叫外卖（半成品自己加热）   IaaS = 自己买菜做饭（给厨房和食材）", 
                0.7, 6.15, 11.9, 0.9, font_size=14, bold=True, color=C_DARK_GRAY)
    
    return slide

def make_nist_slide_5(prs):
    """幻灯片5：5个基本特征"""
    slide = prs.slides.add_slide(blank_layout)
    add_rect(slide, 0, 0, 13.3, 7.5, C_WHITE)
    # 标题栏
    add_rect(slide, 0, 0, 13.3, 1.2, C_GREEN)
    add_textbox(slide, "3-5-4 之【5】五个基本特征", 0.5, 0.2, 12.3, 0.8, font_size=28, bold=True, color=C_WHITE, align=PP_ALIGN.CENTER)
    
    features = [
        ("1", "按需自助服务", "On-demand Self-service", "用户自行获取资源，不需人工干预", "自助餐厅，想拿什么自己拿", C_BLUE, C_LIGHT_BLUE),
        ("2", "广泛网络访问", "Broad Network Access", "通过网络随时随地访问，支持各种终端", "有网就能用，手机电脑都行", C_GREEN, C_LIGHT_GREEN),
        ("3", "资源池化", "Resource Pooling", "多用户共享物理资源，按需动态分配", "所有人共用一个大池子，按需分配", C_ORANGE, C_LIGHT_ORANGE),
        ("4", "快速弹性", "Rapid Elasticity", "资源可快速扩容/缩容，看似无限", "橡皮筋，能大能小能伸能缩", C_RED, C_LIGHT_RED),
        ("5", "可计量服务", "Measured Service", "按使用量计费，资源使用可监控和报告", "水表电表，用多少算多少", C_PURPLE, C_LIGHT_PURPLE),
    ]
    
    for i, (num, cn, en, desc, analogy, color, bg) in enumerate(features):
        y = 1.5 + i * 1.15
        # 编号
        add_shape_with_text(slide, 0.5, y, 0.8, 0.9, num, color, font_size=24)
        # 中文名+英文
        add_rounded_rect(slide, 1.5, y, 3.0, 0.9, bg)
        add_textbox(slide, cn, 1.6, y + 0.05, 2.8, 0.45, font_size=15, bold=True, color=color)
        add_textbox(slide, en, 1.6, y + 0.5, 2.8, 0.35, font_size=10, color=C_DARK_GRAY)
        # 描述
        add_textbox(slide, desc, 4.7, y + 0.05, 4.0, 0.85, font_size=13, color=C_DARK_GRAY)
        # 类比
        add_rounded_rect(slide, 8.9, y, 3.8, 0.9, C_GRAY_BG)
        add_textbox(slide, analogy, 9.0, y + 0.05, 3.6, 0.85, font_size=13, bold=True, color=color)
    
    return slide

def make_nist_slide_6(prs):
    """幻灯片6：4个部署模型"""
    slide = prs.slides.add_slide(blank_layout)
    add_rect(slide, 0, 0, 13.3, 7.5, C_WHITE)
    # 标题栏
    add_rect(slide, 0, 0, 13.3, 1.2, C_ORANGE)
    add_textbox(slide, "3-5-4 之【4】四个部署模型", 0.5, 0.2, 12.3, 0.8, font_size=28, bold=True, color=C_WHITE, align=PP_ALIGN.CENTER)
    
    models = [
        ("公有云", "Public Cloud", "面向公众开放", "阿里云、腾讯云、AWS\n任何人都可以购买使用", C_BLUE, C_LIGHT_BLUE),
        ("私有云", "Private Cloud", "专为单一组织使用", "企业自建云平台\n如银行、政府的内部云", C_GREEN, C_LIGHT_GREEN),
        ("社区云", "Community Cloud", "多个有共同需求的组织共享", "如教育云、医疗云\n多个学校/医院共用", C_PURPLE, C_LIGHT_PURPLE),
        ("混合云", "Hybrid Cloud", "以上两种或多种的组合", "核心数据放私有云\n普通业务放公有云", C_ORANGE, C_LIGHT_ORANGE),
    ]
    
    for i, (cn, en, desc, detail, color, bg) in enumerate(models):
        x = 0.5 + i * 3.15
        # 标题块
        add_shape_with_text(slide, x, 1.5, 2.95, 1.1, cn, color, font_size=22)
        # 英文
        add_textbox(slide, en, x, 2.7, 2.95, 0.4, font_size=13, color=C_DARK_GRAY, align=PP_ALIGN.CENTER)
        # 一句话描述
        add_rounded_rect(slide, x, 3.2, 2.95, 0.8, bg)
        add_textbox(slide, desc, x + 0.1, 3.25, 2.75, 0.7, font_size=13, bold=True, color=color, align=PP_ALIGN.CENTER)
        # 详细说明
        add_textbox(slide, detail, x + 0.1, 4.2, 2.75, 1.5, font_size=13, color=C_DARK_GRAY, align=PP_ALIGN.CENTER)
    
    # 底部对比提示
    add_rounded_rect(slide, 0.5, 6.0, 12.3, 1.1, C_GRAY_BG)
    add_textbox(slide, "核心区别：公有云 = 租别人的房子住    私有云 = 自己建房子自己住    社区云 = 几家人合住一个大院    混合云 = 自己住一套+租一套", 
                0.7, 6.05, 11.9, 1.0, font_size=14, bold=True, color=C_DARK_GRAY)
    
    return slide

def make_nist_slide_7(prs):
    """幻灯片7：完整3-5-4框架总结"""
    slide = prs.slides.add_slide(blank_layout)
    add_rect(slide, 0, 0, 13.3, 7.5, C_DARK)
    # 标题
    add_textbox(slide, "NIST 3-5-4 完整框架", 1.5, 0.3, 10.3, 0.9, font_size=34, bold=True, color=C_WHITE, align=PP_ALIGN.CENTER)
    add_rect(slide, 3, 1.2, 7.3, 0.06, C_ORANGE)
    
    # 左列：3个服务模型
    add_shape_with_text(slide, 0.5, 1.6, 3.8, 0.7, "3 服务模型", C_BLUE, font_size=18)
    items_3 = ["SaaS - 软件即服务", "PaaS - 平台即服务", "IaaS - 基础设施即服务"]
    for i, item in enumerate(items_3):
        add_rounded_rect(slide, 0.5, 2.5 + i * 0.7, 3.8, 0.55, C_DARK_GRAY)
        add_textbox(slide, item, 0.7, 2.53 + i * 0.7, 3.4, 0.5, font_size=14, color=C_WHITE)
    
    # 中列：5个基本特征
    add_shape_with_text(slide, 4.5, 1.6, 4.3, 0.7, "5 基本特征", C_GREEN, font_size=18)
    items_5 = [
        "按需自助服务",
        "广泛网络访问",
        "资源池化",
        "快速弹性",
        "可计量服务",
    ]
    for i, item in enumerate(items_5):
        add_rounded_rect(slide, 4.5, 2.5 + i * 0.55, 4.3, 0.45, C_DARK_GRAY)
        add_textbox(slide, item, 4.7, 2.52 + i * 0.55, 3.9, 0.4, font_size=13, color=C_WHITE)
    
    # 右列：4个部署模型
    add_shape_with_text(slide, 9.2, 1.6, 3.6, 0.7, "4 部署模型", C_ORANGE, font_size=18)
    items_4 = ["公有云", "私有云", "社区云", "混合云"]
    for i, item in enumerate(items_4):
        add_rounded_rect(slide, 9.2, 2.5 + i * 0.7, 3.6, 0.55, C_DARK_GRAY)
        add_textbox(slide, item, 9.4, 2.53 + i * 0.7, 3.2, 0.5, font_size=14, color=C_WHITE)
    
    # 底部映射到教材章节
    add_rounded_rect(slide, 0.5, 5.6, 12.3, 1.5, C_DARK_GRAY)
    add_textbox(slide, "3-5-4 对应教材章节", 0.7, 5.65, 11.9, 0.4, font_size=14, bold=True, color=C_ORANGE)
    add_textbox(slide, "第1章（云计算简介）→ 引出NIST定义    第2章（云计算的服务）→ 3个服务模型", 0.7, 6.05, 11.9, 0.4, font_size=13, color=C_WHITE)
    add_textbox(slide, "第3章（云计算的部署）→ 4个部署模型    第4章（云计算的特点）→ 5个基本特征", 0.7, 6.45, 11.9, 0.4, font_size=13, color=C_WHITE)
    
    return slide

# 生成7张NIST幻灯片
print("  生成 NIST 3-5-4 渐进展示幻灯片（7张）...")
make_nist_slide_1(prs1)  # NIST是什么
make_nist_slide_2(prs1)  # 云计算定义
make_nist_slide_3(prs1)  # 3-5-4框架总览
make_nist_slide_4(prs1)  # 3个服务模型
make_nist_slide_5(prs1)  # 5个基本特征
make_nist_slide_6(prs1)  # 4个部署模型
make_nist_slide_7(prs1)  # 完整框架总结
print("  [OK] 7张NIST幻灯片已添加到第1章PPT末尾")

# 保存
out_path = os.path.join(OUT, '1.云计算的简介.pptx')
prs1.save(out_path)
print(f"  [OK] 保存到 {out_path}")
print(f"  第1章PPT总页数: {len(prs1.slides)} 页（原43页 + 7张NIST = 50页）")

# ============================================================
# 汇总
# ============================================================
print()
print("=" * 60)
print("全部完成！输出文件：")
print("=" * 60)
for fname in ['1.云计算的简介.pptx', '2.云计算的服务.pptx', '3.云计算的部署.pptx']:
    p = os.path.join(OUT, fname)
    prs_check = Presentation(p)
    print(f"  {p}")
    print(f"    总页数: {len(prs_check.slides)}")
