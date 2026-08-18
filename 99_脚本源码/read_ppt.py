# -*- coding: utf-8 -*-
import os, sys
from pptx import Presentation

base = r"E:/xiqin/云计算基础技术与应用-PPT课件.rar/53938-云计算基础技术与应用-PPT课件"
files = ["1.云计算的简介.pptx","2.云计算的服务.pptx","3.云计算的部署.pptx","4.云计算的特点.pptx","5.云计算安全.pptx"]

for fname in files:
    path = os.path.join(base, fname)
    prs = Presentation(path)
    print(f"\n{'='*60}")
    print(f"文件: {fname}  | 幻灯片数: {len(prs.slides)}")
    print(f"{'='*60}")
    for i, slide in enumerate(prs.slides, 1):
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    t = para.text.strip()
                    if t:
                        texts.append(t)
        print(f"\n--- 第{i}页 ---")
        for t in texts:
            print(t)
